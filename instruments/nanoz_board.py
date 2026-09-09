from __future__ import annotations

import csv
import datetime as dt
import json
import os
import re
import shutil
import struct
import threading
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Optional

import serial
from serial.tools import list_ports

BAUD = 921600
READ_TIMEOUT_S = 0.05

VER_RE = re.compile(r"SW:(V[^\s]+).*?S/N:\s*([0-9A-Fa-f]+-[0-9A-Fa-f]+)", re.S)
WHOAMI_RE = re.compile(r"Iam\s+([0-9A-Fa-f]+)", re.I)
ENV_HEADER_RE = re.compile(rb"#env(\d)!\s+(\d+)\s+([0-9A-Fa-f]+)\s+(\d+)\s+(\d+)\s*$", re.I)
SPL_HEADER_RE = re.compile(rb"#spl!\s+(\d+)\s+([0-9A-Fa-f]+)\s+(\d+)\s+(\d+)\s+(\d+)\s*$", re.I)
SEQ_HEADER_RE = re.compile(rb"#seq!\s+(-?\d+)\s+(-?\d+)\s*$", re.I)
EEP_HEADER_RE = re.compile(rb"#eep!\s+(\d+)\s+([0-9A-Fa-f]+)\s+(\d+)\s*$", re.I)


class NanoZError(RuntimeError):
    pass


@dataclass
class PortMeta:
    device: str
    description: str
    hwid: str
    serial_number: str
    vid_pid: str
    location: str


@dataclass
class BoardIdentity:
    port: str
    serial_number: str
    firmware: str
    signature: str
    raw_ver: str
    raw_whoami: str
    usb_id: str
    slot0: Optional[int] = None
    slot1: Optional[int] = None
    last_port: Optional[str] = None

    def chip_slots(self) -> dict:
        return {"0": self.slot0, "1": self.slot1}


def now_stamp() -> str:
    return dt.datetime.now().isoformat(timespec="milliseconds")


def list_serial_ports() -> list[PortMeta]:
    out: list[PortMeta] = []
    for p in sorted(list_ports.comports(), key=lambda x: x.device):
        vid_pid = ""
        if p.vid is not None and p.pid is not None:
            vid_pid = f"{p.vid:04X}:{p.pid:04X}"
        out.append(
            PortMeta(
                device=p.device,
                description=p.description or "",
                hwid=p.hwid or "",
                serial_number=p.serial_number or "",
                vid_pid=vid_pid,
                location=p.location or "",
            )
        )
    return out


def open_serial(port: str) -> serial.Serial:
    ser = serial.Serial(
        port=port,
        baudrate=BAUD,
        bytesize=serial.EIGHTBITS,
        parity=serial.PARITY_NONE,
        stopbits=serial.STOPBITS_ONE,
        timeout=READ_TIMEOUT_S,
        write_timeout=1.0,
        xonxoff=False,
        rtscts=False,
        dsrdtr=False,
    )
    time.sleep(0.2)
    ser.reset_input_buffer()
    ser.reset_output_buffer()
    return ser


def send_ascii(ser: serial.Serial, cmd: str) -> None:
    if not cmd.endswith("\r"):
        cmd += "\r"
    ser.write(cmd.encode("ascii"))
    ser.flush()


def read_text_for(ser: serial.Serial, seconds: float) -> str:
    deadline = time.time() + seconds
    data = bytearray()
    while time.time() < deadline:
        chunk = ser.read(4096)
        if chunk:
            data.extend(chunk)
        else:
            time.sleep(0.01)
    return data.decode(errors="replace").strip()


def identify_on_port(port: str) -> Optional[BoardIdentity]:
    ser = open_serial(port)
    try:
        send_ascii(ser, "ver")
        raw_ver = read_text_for(ser, 0.75)
        send_ascii(ser, "whoami")
        raw_whoami = read_text_for(ser, 0.50)
    except Exception:
        return None
    finally:
        ser.close()

    m = VER_RE.search(raw_ver)
    if not m:
        return None

    firmware, sn = m.group(1), m.group(2).upper()
    wm = WHOAMI_RE.search(raw_whoami)
    signature = wm.group(1) if wm else ""

    usb_id = ""
    for p in list_serial_ports():
        if p.device.upper() == port.upper():
            usb_id = p.serial_number or f"VIDPID={p.vid_pid};LOC={p.location};PORT={p.device}"
            break

    return BoardIdentity(
        port=port,
        serial_number=sn,
        firmware=firmware,
        signature=signature,
        raw_ver=raw_ver,
        raw_whoami=raw_whoami,
        usb_id=usb_id,
    )


def discover_boards(ports: Optional[list[str]] = None,
                    log: Optional[Callable[[str], None]] = None) -> list[BoardIdentity]:
    candidates = ports if ports is not None else [p.device for p in list_serial_ports()]
    found: list[BoardIdentity] = []
    for port in candidates:
        if log:
            log(f"Probing {port}...")
        try:
            ident = identify_on_port(port)
        except Exception as e:
            if log:
                log(f"  -> could not open ({e}) - in use by another program?")
            continue
        if ident:
            found.append(ident)
            if log:
                log(f"  -> NanoZ board found: S/N {ident.serial_number}  FW {ident.firmware}")
        elif log:
            log(f"  -> no response (not a NanoZ board, or powered off)")
    return found


def read_line_bytes(ser: serial.Serial, buffer: bytearray, timeout_s: float = 2.0) -> Optional[bytes]:
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        idx = buffer.find(b"\n")
        if idx >= 0:
            line = bytes(buffer[:idx]).strip(b"\r\n ")
            del buffer[: idx + 1]
            if line:
                return line
            continue
        chunk = ser.read(4096)
        if chunk:
            buffer.extend(chunk)
        else:
            time.sleep(0.005)
    return None


def read_exact_from_buffer(ser: serial.Serial, buffer: bytearray, n: int, timeout_s: float = 2.0) -> bytes:
    deadline = time.time() + timeout_s
    while len(buffer) < n and time.time() < deadline:
        chunk = ser.read(n - len(buffer))
        if chunk:
            buffer.extend(chunk)
        else:
            time.sleep(0.005)
    if len(buffer) < n:
        raise NanoZError(f"Timed out waiting for binary block: needed {n}, got {len(buffer)}")
    data = bytes(buffer[:n])
    del buffer[:n]
    return data


def parse_spl_data(data: bytes) -> dict:
    if len(data) < 48:
        raise NanoZError(f"SPL data block too short: {len(data)} bytes")
    vals = struct.unpack_from("<IBBH4h4f4f", data, 0)
    ppms, chip_id, sensor_mask, reserved = vals[:4]
    dac = vals[4:8]
    adc = vals[8:12]
    heaters = vals[12:16]
    return {
        "ppms": ppms,
        "chip_id": chip_id,
        "sensor_mask": sensor_mask,
        "reserved": reserved,
        "dac_mv_s1": dac[0],
        "dac_mv_s2": dac[1],
        "dac_mv_s3": dac[2],
        "dac_mv_s4": dac[3],
        "adc_current_ma_s1": adc[0],
        "adc_current_ma_s2": adc[1],
        "adc_current_ma_s3": adc[2],
        "adc_current_ma_s4": adc[3],
        "heater1_voltage_mv": heaters[0],
        "heater1_current_ma": heaters[1],
        "heater2_voltage_mv": heaters[2],
        "heater2_current_ma": heaters[3],
    }


def parse_env_data(data: bytes) -> dict:
    if len(data) < 132:
        raise NanoZError(f"ENV data block too short: {len(data)} bytes")
    off = 0
    pps, samples_nb, adc_mask = struct.unpack_from("<IHH", data, off)
    off += 8
    adc_samples = struct.unpack_from("<8H", data, off)
    off += 16
    adc_voltage = struct.unpack_from("<8f", data, off)
    off += 32
    adc_current = struct.unpack_from("<8f", data, off)
    off += 32
    htr_voltage = struct.unpack_from("<4f", data, off)
    off += 16
    adc_mid, mcu_temp = struct.unpack_from("<ff", data, off)
    off += 8
    humidity_x100, tempH_x100, pressure_x10, tempP_x100, pending, align = struct.unpack_from("<6h", data, off)
    off += 12
    age = struct.unpack_from("<2I", data, off)

    return {
        "pps": pps,
        "adc_samples_nb": samples_nb,
        "adc_mask": adc_mask,
        "adc_mid_value": adc_mid,
        "mcu_temperature_c": mcu_temp,
        "humidity_percent": humidity_x100 / 100.0,
        "temp_h_c": tempH_x100 / 100.0,
        "pressure_hpa_minus_1013": pressure_x10 / 10.0,
        "temp_p_c": tempP_x100 / 100.0,
        "pending": pending,
        "align": align,
        "age_chip1_s": age[0],
        "age_chip2_s": age[1],
        "adc_samples_4x2": ";".join(str(x) for x in adc_samples),
        "adc_voltage_4x2": ";".join(f"{x:.6g}" for x in adc_voltage),
        "adc_current_4x2": ";".join(f"{x:.6g}" for x in adc_current),
        "htr_voltage_2x2": ";".join(f"{x:.6g}" for x in htr_voltage),
    }


EEPROM_PAGE_SIZE = 32
EEPROM_PARAMS_ADDR = 0
EEPROM_CYCLES_PAGE = 16
EEPROM_CYCLES_ADDR = EEPROM_CYCLES_PAGE * EEPROM_PAGE_SIZE
EEPROM_SEQUENCES_PAGE = 64
EEPROM_SEQUENCES_ADDR = EEPROM_SEQUENCES_PAGE * EEPROM_PAGE_SIZE
EEPROM_CYCLE_RECORD_SIZE = 32
MAX_CYCLES_NB = 48
MAX_SEQUENCE_NB = 96


def parse_params_block(data: bytes) -> dict:
    if len(data) < 152:
        raise NanoZError(f"PARAMS block too short: {len(data)} bytes (need >= 152)")
    signature = struct.unpack_from("<H", data, 0)[0]
    cycles_configured = struct.unpack_from("<H", data, 4)[0]
    periodicity_ms = data[14]
    cal1, cal2 = struct.unpack_from("<ff", data, 20)

    def chip_record(off):
        w, x, y, z, age_s = struct.unpack_from("<5I", data, off)
        return {"w": w, "x": x, "y": y, "z": z, "age_s": age_s,
               "id": f"D{w}L{x}-{y}-{z}"}

    return {
        "signature": f"0x{signature:04X}",
        "cycles_configured": cycles_configured,
        "periodicity_ms": periodicity_ms,
        "cal1": cal1,
        "cal2": cal2,
        "chip1": chip_record(112),
        "chip2": chip_record(132),
    }


def parse_cycle_record(data: bytes) -> "dict | None":
    if len(data) < EEPROM_CYCLE_RECORD_SIZE:
        raise NanoZError(f"Cycle record too short: {len(data)} bytes")
    if all(b == 0xFF for b in data[:EEPROM_CYCLE_RECORD_SIZE]):
        return None
    wire_index, num_sequences = struct.unpack_from("<HH", data, 0)
    seq_refs = []
    off = 4
    for _ in range(min(num_sequences, (EEPROM_CYCLE_RECORD_SIZE - 4) // 4)):
        seq_refs.append(struct.unpack_from("<I", data, off)[0])
        off += 4
    return {"wire_index": wire_index, "num_sequences": num_sequences,
           "sequence_refs": seq_refs}


SEQ_FIELD_OFFSETS = {
    "duration_s": 2,
    "delay_s": 4,
    "sensor_mv": 6,
    "ramp_up_ms": 22,
    "high_duration_ms": 26,
    "ramp_down_ms": 30,
    "low_duration_ms": 34,
    "phase_shift_ms": 38,
    "heater1_low_mv": 42,
    "heater2_low_mv": 46,
    "heater1_high_mv": 50,
    "heater2_high_mv": 54,
    "chip": 58,
    "resolution_ms": 62,
}


def parse_sequence_records(data: bytes) -> list:
    records = []
    start = 0
    n = len(data)
    while start < n:
        if data[start] == 0xFF and (start + 1 >= n or data[start + 1] == 0xFF):
            break
        term = data.find(b"\xff\xff", start)
        if term == -1:
            end = n
        else:
            end = term + 2
        record = data[start:end]
        if len(record) >= 4:
            wire_index = struct.unpack_from("<H", record, 0)[0]
            duration_s = struct.unpack_from("<h", record, 2)[0]
            delay_s = struct.unpack_from("<h", record, 4)[0] if len(record) >= 6 else None
            sensors_mv = sensors_pad_raw = None
            if len(record) >= 22:
                sensors_mv = [struct.unpack_from("<h", record, o)[0] for o in (6, 10, 14, 18)]
                sensors_pad_raw = [struct.unpack_from("<h", record, o)[0] for o in (8, 12, 16, 20)]
            heater_times_ms = None
            if len(record) >= 32:
                heater_times_ms = [struct.unpack_from("<h", record, o)[0] for o in (22, 26, 30)]
            heater_extra = None
            if len(record) >= 40:
                heater_extra = [struct.unpack_from("<h", record, o)[0] for o in (34, 38)]
            heater_voltages_mv = heater_v_pad_raw = None
            if len(record) >= 58:
                heater_voltages_mv = [struct.unpack_from("<h", record, o)[0] for o in (42, 46, 50, 54)]
                heater_v_pad_raw = [struct.unpack_from("<h", record, o)[0] for o in (44, 48, 52, 56)]
            chip_candidates = None
            if len(record) >= 62:
                chip_candidates = [struct.unpack_from("<h", record, o)[0] for o in (58, 60)]
            resolution_candidates = None
            if len(record) >= 66:
                resolution_candidates = [struct.unpack_from("<h", record, o)[0] for o in (62, 64)]
            ht = heater_times_ms or [None, None, None]
            he = heater_extra or [None, None]
            hv = heater_voltages_mv or [None, None, None, None]
            records.append({
                "blob_offset": start,
                "record_len": len(record),
                "wire_index": wire_index,
                "duration_s": duration_s,
                "delay_s": delay_s,
                "sensor_mv": sensors_mv[0] if sensors_mv else None,
                "sensors_mv": sensors_mv,
                "sensors_pad_raw": sensors_pad_raw,
                "ramp_up_ms": ht[0],
                "high_duration_ms": ht[1],
                "ramp_down_ms": ht[2],
                "low_duration_ms": he[0],
                "phase_shift_ms": he[1],
                "heater1_low_mv": hv[0],
                "heater2_low_mv": hv[1],
                "heater1_high_mv": hv[2],
                "heater2_high_mv": hv[3],
                "heater_times_ms": heater_times_ms,
                "heater_extra": heater_extra,
                "heater_voltages_mv": heater_voltages_mv,
                "heater_v_pad_raw": heater_v_pad_raw,
                "chip_candidates": chip_candidates,
                "resolution_candidates": resolution_candidates,
                "chip": chip_candidates[0] if chip_candidates else None,
                "resolution_ms": resolution_candidates[0] if resolution_candidates else None,
                "raw_hex": record.hex(),
            })
        start = end
    return records


def encode_sequence_patch(original_record: bytes, fields: dict) -> bytearray:
    buf = bytearray(original_record)

    def put_u(offset, value):
        struct.pack_into("<H", buf, offset, int(value))

    def put_s(offset, value):
        struct.pack_into("<h", buf, offset, int(value))

    if "duration_s" in fields:
        put_u(SEQ_FIELD_OFFSETS["duration_s"], fields["duration_s"])
    if "delay_s" in fields:
        put_u(SEQ_FIELD_OFFSETS["delay_s"], fields["delay_s"])
    if "sensor_mv" in fields:
        for o in (6, 10, 14, 18):
            put_s(o, fields["sensor_mv"])
    for key in ("ramp_up_ms", "high_duration_ms", "ramp_down_ms", "low_duration_ms",
               "phase_shift_ms", "heater1_low_mv", "heater2_low_mv",
               "heater1_high_mv", "heater2_high_mv"):
        if key in fields:
            put_u(SEQ_FIELD_OFFSETS[key], fields[key])
    return buf


def append_csv_row(path, row: dict) -> None:
    path = Path(path)
    exists = path.exists() and path.stat().st_size > 0
    with path.open("a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(row.keys()))
        if not exists:
            writer.writeheader()
        writer.writerow(row)


BOARDS_MEMORY_FILENAME = "ata_nanoz_boards.json"


def save_known_boards(folder, identities: list) -> None:
    by_sn: dict[str, dict] = {}
    for i in identities:
        by_sn[i.serial_number or f"(no S/N) {i.port}"] = {
            "serial_number": i.serial_number, "firmware": i.firmware,
            "signature": i.signature, "usb_id": i.usb_id, "slot0": i.slot0, "slot1": i.slot1,
            "last_port": i.port or i.last_port or None,
        }
    path = Path(folder) / BOARDS_MEMORY_FILENAME
    path.write_text(json.dumps(list(by_sn.values()), indent=2), encoding="utf-8")


def load_known_boards(folder) -> list[BoardIdentity]:
    path = Path(folder) / BOARDS_MEMORY_FILENAME
    if not path.is_file():
        return []
    try:
        rows = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []
    return [
        BoardIdentity(
            port="", serial_number=row.get("serial_number", ""),
            firmware=row.get("firmware", ""), signature=row.get("signature", ""),
            raw_ver="", raw_whoami="", usb_id=row.get("usb_id", ""),
            slot0=row.get("slot0", row.get("slot")),
            slot1=row.get("slot1"),
            last_port=row.get("last_port"),
        )
        for row in rows
    ]


PROBE_HEIGHT_FILENAME = "ata_nanoz_probe_height.json"


def save_probe_height(folder, n: int) -> None:
    path = Path(folder) / PROBE_HEIGHT_FILENAME
    path.write_text(json.dumps({"probe_height": int(n)}), encoding="utf-8")


def load_probe_height(folder) -> int:
    path = Path(folder) / PROBE_HEIGHT_FILENAME
    if not path.is_file():
        return DEFAULT_PROBE_HEIGHT
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        n = int(data.get("probe_height", DEFAULT_PROBE_HEIGHT))
        return n if n > 0 else DEFAULT_PROBE_HEIGHT
    except (OSError, ValueError, TypeError):
        return DEFAULT_PROBE_HEIGHT


WAFER_PLAN_XLSX_FILENAME = "ata_nanoz_wafer_plan.xlsx"


def wafer_plan_path_in_folder(folder) -> str:
    return str(Path(folder) / WAFER_PLAN_XLSX_FILENAME)


def import_wafer_plan_into_folder(folder, source_path: str) -> str:
    dest = wafer_plan_path_in_folder(folder)
    if os.path.abspath(source_path) != os.path.abspath(dest):
        shutil.copyfile(source_path, dest)
    return dest


LEGACY_RECIPE_FILENAME = "ata_nanoz_recipe.json"
RECIPES_FILENAME = "ata_nanoz_recipes.json"

_RECIPE_SHOT_META_KEYS = ("die_column", "td_start_row", "td_end_row", "board_reasons",
                          "chip_reasons")


def _shots_to_rows(shots: list) -> list:
    rows = []
    for s in shots:
        row = {"label": s.get("label", ""), "excluded_boards": sorted(s.get("excluded_boards", ()))}
        for k in _RECIPE_SHOT_META_KEYS:
            if k in s:
                row[k] = s[k]
        rows.append(row)
    return rows


def _rows_to_shots(rows: list) -> list[dict]:
    shots = []
    for s in rows:
        if not isinstance(s, dict):
            continue
        shot = {"label": s.get("label", ""), "excluded_boards": set(s.get("excluded_boards", ()))}
        for k in _RECIPE_SHOT_META_KEYS:
            if k in s:
                shot[k] = s[k]
        shots.append(shot)
    return shots


def _load_recipes_file(folder) -> dict:
    path = Path(folder) / RECIPES_FILENAME
    if path.is_file():
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(data.get("recipes"), dict):
                data.setdefault("active", None)
                return data
        except (OSError, ValueError):
            pass
    return {"active": None, "recipes": {}}


def _write_recipes_file(folder, data: dict) -> None:
    path = Path(folder) / RECIPES_FILENAME
    path.write_text(json.dumps(data, indent=2), encoding="utf-8")


def list_recipe_names(folder) -> list[str]:
    return sorted(_load_recipes_file(folder)["recipes"].keys())


def get_active_recipe_name(folder):
    return _load_recipes_file(folder).get("active")


def save_named_recipe(folder, name: str, shots: list, wafer_plan_path: str | None = None,
                      touchdowns: "list | None" = None) -> None:
    data = _load_recipes_file(folder)
    data["recipes"][name] = _shots_to_rows(shots)
    data["active"] = name
    if wafer_plan_path:
        data.setdefault("wafer_plan_paths", {})[name] = wafer_plan_path
    if touchdowns is not None:
        data.setdefault("touchdowns", {})[name] = touchdowns
    _write_recipes_file(folder, data)


def load_named_recipe(folder, name: str) -> list[dict]:
    rows = _load_recipes_file(folder)["recipes"].get(name)
    return _rows_to_shots(rows) if rows is not None else []


def load_named_touchdowns(folder, name: str) -> list[dict]:
    rows = _load_recipes_file(folder).get("touchdowns", {}).get(name)
    return list(rows) if isinstance(rows, list) else []


def get_recipe_wafer_plan_path(folder, name: str) -> str | None:
    return _load_recipes_file(folder).get("wafer_plan_paths", {}).get(name)


def set_active_recipe(folder, name: str) -> None:
    data = _load_recipes_file(folder)
    if name in data["recipes"]:
        data["active"] = name
        _write_recipes_file(folder, data)


def delete_named_recipe(folder, name: str) -> None:
    data = _load_recipes_file(folder)
    data["recipes"].pop(name, None)
    data.get("wafer_plan_paths", {}).pop(name, None)
    data.get("touchdowns", {}).pop(name, None)
    if data.get("active") == name:
        data["active"] = None
    _write_recipes_file(folder, data)


def load_active_recipe(folder):
    data = _load_recipes_file(folder)
    name = data.get("active")
    if not name or name not in data["recipes"]:
        return None, [], None
    wafer_plan_path = data.get("wafer_plan_paths", {}).get(name)
    return name, _rows_to_shots(data["recipes"][name]), wafer_plan_path


def migrate_legacy_recipe(folder):
    legacy_path = Path(folder) / LEGACY_RECIPE_FILENAME
    if not legacy_path.is_file() or _load_recipes_file(folder)["recipes"]:
        return None
    try:
        legacy = json.loads(legacy_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    shots = _rows_to_shots(legacy.get("shots", []))
    if not shots:
        return None
    save_named_recipe(folder, "Imported", shots)
    return "Imported"


try:
    import openpyxl
    _OPENPYXL_AVAILABLE = True
except ImportError:
    _OPENPYXL_AVAILABLE = False

_REFERENCE_FILL_RGBS = frozenset({"FFC00000"})

DEFAULT_PROBE_HEIGHT = 20


@dataclass
class WaferPlan:
    dies: dict
    serial_to_rc: dict
    touchdowns: list
    probe_height: int = DEFAULT_PROBE_HEIGHT


def load_wafer_plan(path, probe_height: int = DEFAULT_PROBE_HEIGHT) -> WaferPlan:
    if not _OPENPYXL_AVAILABLE:
        raise NanoZError("openpyxl is required to import a wafer plan .xlsx (pip install openpyxl)")

    wb = openpyxl.load_workbook(path, data_only=True)
    for name in ("Die Map", "Touchdown List"):
        if name not in wb.sheetnames:
            raise NanoZError(f"'{name}' sheet not found — not a recognized wafer-plan workbook.")

    die_ws = wb["Die Map"]
    dies: dict[tuple[int, int], dict] = {}
    serial_to_rc: dict[str, tuple[int, int]] = {}
    for row in die_ws.iter_rows(min_row=3, min_col=2):
        for cell in row:
            if not cell.value:
                continue
            serial = str(cell.value).strip()
            r, c = cell.row - 2, cell.column - 1
            fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else None
            status = "reference" if fill in _REFERENCE_FILL_RGBS else "product"
            dies[(r, c)] = {"serial": serial, "status": status}
            serial_to_rc[serial.upper()] = (r, c)
    if not dies:
        raise NanoZError("Die Map: no dies found.")

    td_ws = wb["Touchdown List"]
    touchdowns = []
    missing = []
    for row in td_ws.iter_rows(min_row=3, max_col=1, values_only=True):
        serial = row[0] if row else None
        if not serial:
            continue
        serial = str(serial).strip()
        rc = serial_to_rc.get(serial.upper())
        if rc is None:
            missing.append(serial)
            continue
        touchdowns.append(rc)
    if not touchdowns:
        raise NanoZError("Touchdown List: no touchdown dies found.")
    if missing:
        raise NanoZError(
            f"Touchdown List references {len(missing)} die ID(s) not found on Die Map: "
            + ", ".join(missing[:5]) + (", ..." if len(missing) > 5 else ""))

    return WaferPlan(dies=dies, serial_to_rc=serial_to_rc, touchdowns=touchdowns,
                     probe_height=probe_height)


def tile_windows_covering_wafer(die_keys, window_height: int) -> list:
    window_height = max(1, int(window_height or 1))
    rows_by_col: dict = {}
    for r, c in die_keys:
        rows_by_col.setdefault(c, []).append(r)
    picks = []
    for c in sorted(rows_by_col):
        rows = sorted(rows_by_col[c])
        rows_set = set(rows)
        start, max_r = rows[0], rows[-1]
        while start <= max_r:
            end = start + window_height - 1
            top = next((r for r in range(start, end + 1) if r in rows_set), None)
            if top is not None:
                picks.append((top, c))
            start += window_height
    return picks


def classify_die(plan: "WaferPlan", row: int, col: int,
                 row_offset: int = 0, col_offset: int = 0) -> str:
    d = plan.dies.get((row - row_offset, col - col_offset))
    return d["status"] if d else "off_wafer"


def touchdown_slot_exclusions(die_col: int, start_row: int, end_row: int, plan: "WaferPlan",
                              row_offset: int = 0, col_offset: int = 0) -> dict:
    result = {}
    for slot in range(1, plan.probe_height + 1):
        physical_row = start_row + slot - 1
        if physical_row > end_row:
            result[slot] = "past touchdown end"
            continue
        status = classify_die(plan, physical_row, die_col, row_offset, col_offset)
        result[slot] = {"off_wafer": "off wafer", "reference": "reference die",
                        "product": None}[status]
    return result


def wafer_plan_die_grid(plan: "WaferPlan") -> list[dict]:
    return [{"row": r, "col": c, "status": d["status"], "serial": d["serial"]}
           for (r, c), d in sorted(plan.dies.items())]


def wafer_plan_stats(plan: "WaferPlan") -> dict:
    counts = {"product": 0, "reference": 0, "off_wafer": 0}
    for start_row, die_col in plan.touchdowns:
        end_row = start_row + plan.probe_height - 1
        for reason in touchdown_slot_exclusions(die_col, start_row, end_row, plan).values():
            if reason is None:
                counts["product"] += 1
            elif reason == "reference die":
                counts["reference"] += 1
            else:
                counts["off_wafer"] += 1
    return counts


def _build_shot(plan: "WaferPlan", die_col: int, start: int, end: int, ports: list,
                slots_by_port: dict, label: str,
                row_offset: int = 0, col_offset: int = 0,
                already_covered: "set | None" = None) -> dict:
    exclusions = touchdown_slot_exclusions(die_col, start, end, plan, row_offset, col_offset)
    excluded_boards = set()
    board_reasons = {}
    chip_reasons = {}
    newly_covered = set()
    for port in ports:
        chip_slots = slots_by_port.get(port) or {}
        per_chip = {}
        for chip in ("0", "1"):
            slot = chip_slots.get(chip)
            if slot is None:
                per_chip[chip] = "no slot assigned"
                continue
            reason = exclusions.get(slot, "slot beyond probe head height")
            if reason is None and already_covered is not None:
                rc = (start + slot - 1, die_col)
                if rc in already_covered:
                    reason = "already probed by an earlier touchdown in this recipe"
                else:
                    newly_covered.add(rc)
            per_chip[chip] = reason
        chip_reasons[port] = per_chip
        if all(r is not None for r in per_chip.values()):
            excluded_boards.add(port)
            board_reasons[port] = "; ".join(
                f"chip{c}: {r}" for c, r in per_chip.items())
        else:
            board_reasons[port] = None
    if already_covered is not None:
        already_covered |= newly_covered
    return {
        "label": label,
        "excluded_boards": excluded_boards,
        "board_reasons": board_reasons,
        "chip_reasons": chip_reasons,
        "die_column": die_col, "td_start_row": start, "td_end_row": end,
    }


def build_shots_from_windows(plan: "WaferPlan", windows: list, ports: list,
                             slots_by_port: dict,
                             row_offset: int = 0, col_offset: int = 0) -> list[dict]:
    shots = []
    covered: set = set()
    for start_row, die_col in windows:
        end = start_row + plan.probe_height - 1
        d = plan.dies.get((start_row - row_offset, die_col - col_offset))
        label = d["serial"] if d else f"Col {die_col}, Row {start_row}"
        shots.append(_build_shot(plan, die_col, start_row, end, ports, slots_by_port, label,
                                 row_offset, col_offset, already_covered=covered))
    return shots


def active_ports_for_window(plan: "WaferPlan", die_col: int, start_row: int,
                            ports: list, slots_by_port: dict,
                            row_offset: int = 0, col_offset: int = 0) -> list:
    end_row = start_row + plan.probe_height - 1
    shot = _build_shot(plan, die_col, start_row, end_row, ports, slots_by_port, "",
                       row_offset, col_offset)
    return [p for p in ports if p not in shot["excluded_boards"]]


class NanoZBoard:

    def __init__(self, identity: BoardIdentity, out_queue,
                die_provider: Optional[Callable[[Optional[str]], tuple]] = None,
                env_interval_s: float = 1.0):
        self.identity = identity
        self.port = identity.port
        self.out_queue = out_queue
        self._die_provider = die_provider or (lambda chip: (None, None, None))
        self._active_die: dict = {}
        self.env_interval_s = env_interval_s
        self.ser: Optional[serial.Serial] = None
        self._buffer = bytearray()
        self._thread: Optional[threading.Thread] = None
        self._running = False
        self.spl_count = 0
        self.env_count = 0
        self.last_error = ""

    def _die_for(self, chip_key) -> tuple:
        if chip_key in self._active_die:
            return self._active_die[chip_key]
        return self._die_provider(chip_key)

    def set_active_die(self, die_map: dict):
        self._active_die = dict(die_map)

    @property
    def is_running(self) -> bool:
        return self._running

    @property
    def state(self) -> str:
        if self._running and self.last_error:
            return "error"
        if self._running:
            return "connected"
        return "not_connected"

    def connect(self):
        if self.ser is None:
            self.ser = open_serial(self.port)

    def start(self):
        self.last_error = ""
        self.connect()
        try:
            send_ascii(self.ser, "pause")
            time.sleep(0.2)
            self.ser.reset_input_buffer()
        except Exception:
            pass
        self._running = True
        self._thread = threading.Thread(target=self._reader_loop, daemon=True)
        self._thread.start()

    def reconnect(self):
        self.stop()
        self.start()

    def run_cycle(self, cycle: int):
        if self.ser:
            send_ascii(self.ser, f"run {cycle}")

    def pause(self):
        if self.ser:
            send_ascii(self.ser, "pause")

    def request_eeprom(self, addr: int, length: int):
        if self.ser:
            send_ascii(self.ser, f"rdeep {addr} {length}")

    def send_raw(self, cmd: str):
        if self.ser:
            send_ascii(self.ser, cmd)

    def write_eeprom(self, addr: int, data: bytes):
        if not self.ser:
            return
        cs = 0
        for b in data:
            cs ^= b
        send_ascii(self.ser, f"wreep {len(data)} {cs:04X} {addr}")
        self.ser.write(bytes(data))
        self.ser.flush()

    def stop(self):
        self._running = False
        if self._thread:
            self._thread.join(timeout=2.0)
            self._thread = None
        if self.ser:
            try:
                self.ser.close()
            except Exception:
                pass
        self.ser = None


    def _reader_loop(self):
        next_env = time.time() + self.env_interval_s if self.env_interval_s > 0 else float("inf")
        while self._running:
            try:
                if self.env_interval_s > 0 and time.time() >= next_env:
                    send_ascii(self.ser, "#env?")
                    next_env = time.time() + self.env_interval_s
                line = read_line_bytes(self.ser, self._buffer, timeout_s=0.2)
            except Exception as e:
                self.last_error = str(e)
                time.sleep(0.2)
                continue
            if line is None:
                continue
            self.last_error = ""
            if not line.startswith(b"#") or line.startswith(b"##"):
                self._emit_text(line)
                continue
            sm = SPL_HEADER_RE.match(line)
            if sm:
                self._handle_spl(sm)
                continue
            em = ENV_HEADER_RE.match(line)
            if em:
                self._handle_env(em)
                continue
            eepm = EEP_HEADER_RE.match(line)
            if eepm:
                self._handle_eep(eepm)
                continue
            if SEQ_HEADER_RE.match(line):
                self._emit_text(line)
                continue
            self.out_queue.put({
                "kind": "unrecognized", "board_sn": self.identity.serial_number,
                "port": self.port, "raw": line, "host_timestamp": now_stamp(),
            })

    def _emit_text(self, line: bytes):
        self.out_queue.put({
            "kind": "text", "board_sn": self.identity.serial_number,
            "port": self.port, "text": line.decode(errors="replace"),
            "host_timestamp": now_stamp(),
        })

    def _handle_spl(self, m):
        length_s, cs_s, chip_s, time_s, bfr_s = m.groups()
        length, header_chip, header_time, header_bfr = (
            int(length_s), int(chip_s), int(time_s), int(bfr_s))
        expected_cs = int(cs_s, 16)
        try:
            data = read_exact_from_buffer(self.ser, self._buffer, length, timeout_s=2.0)
        except NanoZError as e:
            self.last_error = str(e)
            return
        try:
            parsed = parse_spl_data(data)
        except Exception as e:
            parsed = {"parse_error": str(e)}
        row, col, die_id = self._die_for(str(header_chip))
        self.spl_count += 1
        self.out_queue.put({
            "kind": "spl", "host_timestamp": now_stamp(),
            "board_sn": self.identity.serial_number, "port": self.port,
            "die_row": row, "die_col": col, "die_id": die_id,
            "header_chip": header_chip, "header_time_ms": header_time,
            "header_bfr": header_bfr, "len": length,
            "checksum_expected": expected_cs,
            **parsed,
        })

    def _handle_env(self, m):
        x_s, length_s, cs_s, time_s, bfr_s = m.groups()
        env_x, length, header_time, header_bfr = (
            int(x_s), int(length_s), int(time_s), int(bfr_s))
        expected_cs = int(cs_s, 16)
        try:
            data = read_exact_from_buffer(self.ser, self._buffer, length, timeout_s=2.0)
        except NanoZError as e:
            self.last_error = str(e)
            return
        try:
            parsed = parse_env_data(data)
        except Exception as e:
            parsed = {"parse_error": str(e)}
        row, col, die_id = self._die_for(None)
        self.env_count += 1
        self.out_queue.put({
            "kind": "env", "host_timestamp": now_stamp(),
            "board_sn": self.identity.serial_number, "port": self.port,
            "die_row": row, "die_col": col, "die_id": die_id,
            "env_x": env_x, "header_time_ms": header_time, "header_bfr": header_bfr,
            "len": length, "checksum_expected": expected_cs,
            **parsed,
        })

    def _handle_eep(self, m):
        length_s, cs_s, addr_s = m.groups()
        length, addr = int(length_s), int(addr_s)
        expected_cs = int(cs_s, 16)
        try:
            data = read_exact_from_buffer(self.ser, self._buffer, length, timeout_s=2.0)
        except NanoZError as e:
            self.last_error = str(e)
            return
        actual_cs = 0
        for b in data:
            actual_cs ^= b
        self.out_queue.put({
            "kind": "eep", "host_timestamp": now_stamp(),
            "board_sn": self.identity.serial_number, "port": self.port,
            "addr": addr, "len": length,
            "checksum_expected": expected_cs, "checksum_actual": actual_cs,
            "checksum_ok": actual_cs == expected_cs,
            "data_hex": data.hex(),
        })
